"""Deterministic ALL ALS data-dictionary and PHI/PII header preflight."""

import csv
from pathlib import Path
import re

import yaml


def normalize(value):
    """Conservative comparison key for views and column names."""
    value = str(value or '').lstrip('\ufeff').strip().casefold()
    return re.sub(r'[\s_-]+', '_', value).strip('_')


def _unique(values, label):
    result = []
    seen = set()
    for value in values or []:
        key = normalize(value)
        if not key:
            raise ValueError(f'{label} cannot contain an empty value')
        if key in seen:
            raise ValueError(f'{label} contains duplicate normalized value {key!r}')
        seen.add(key)
        result.append(key)
    return result


def read_headers(path):
    """Read tabular headers without reading/logging data values."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in ('.csv', '.tsv', '.txt'):
        with open(path, newline='', encoding='utf-8-sig') as stream:
            headers = next(csv.reader(stream, delimiter='\t' if suffix == '.tsv' else ','), [])
    elif suffix in ('.xlsx', '.xls'):
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1), [])]
            workbook.close()
        except ImportError as error:
            raise ValueError(f'Cannot inspect Excel headers in {path}: openpyxl is required') from error
    else:
        raise ValueError(f'Unsupported ALL ALS preflight file type: {path.name}')
    normalized = _unique(headers, f'headers in {path.name}')
    if not normalized:
        raise ValueError(f'No headers found in {path.name}')
    return normalized


def load_dictionary(path, *, view_column='View Name', field_column='Field'):
    """Parse sparse or fully populated dictionary rows into view -> field set."""
    rows = _read_dictionary_rows(path)
    if not rows:
        raise ValueError('ALL ALS data dictionary is empty')
    required = {view_column, field_column}
    absent = required - set(rows[0])
    if absent:
        raise ValueError(f'ALL ALS data dictionary missing required columns: {sorted(absent)}')
    views, current_view = {}, None
    for number, row in enumerate(rows, 2):
        raw_view = str(row.get(view_column) or '').strip()
        raw_field = str(row.get(field_column) or '').strip()
        if raw_view:
            current_view = normalize(raw_view)
            if not current_view:
                raise ValueError(f'Dictionary row {number}: empty normalized view name')
            if current_view in views and views[current_view]['display_name'] != raw_view:
                raise ValueError(f'Dictionary row {number}: conflicting spellings for view {current_view!r}')
            views.setdefault(current_view, {'display_name': raw_view, 'fields': set()})
        if not current_view:
            raise ValueError(f'Dictionary row {number}: field appears before a view name')
        # Rows that describe permissible values commonly leave Field empty; they
        # belong to the preceding field but do not represent an input column.
        if not raw_field:
            continue
        field = normalize(raw_field)
        if not field:
            raise ValueError(f'Dictionary row {number}: empty normalized field name')
        if field in views[current_view]['fields']:
            raise ValueError(f'Dictionary row {number}: duplicate field {raw_field!r} in {views[current_view]["display_name"]!r}')
        views[current_view]['fields'].add(field)
    if not views or any(not value['fields'] for value in views.values()):
        raise ValueError('Every ALL ALS dictionary view must contain at least one field')
    return views


def _read_dictionary_rows(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as stream:
            return list(csv.DictReader(stream))
    if suffix in ('.xlsx', '.xls'):
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows, None)
            result = [dict(zip(headers, row)) for row in rows]
            workbook.close()
            return result
        except ImportError as error:
            raise ValueError(f'Cannot inspect Excel dictionary {path}: openpyxl is required') from error
    raise ValueError(f'Unsupported ALL ALS data dictionary type: {path.name}')


def load_policy(path):
    try:
        with open(path) as stream:
            policy = yaml.safe_load(stream) or {}
    except FileNotFoundError as error:
        raise ValueError(f'ALL ALS PHI/PII policy file not found: {path}') from error
    if policy.get('version') != 1 or not isinstance(policy.get('forbidden_columns'), list):
        raise ValueError('ALL ALS PHI/PII policy must have version: 1 and forbidden_columns list')
    return set(_unique(policy['forbidden_columns'], 'forbidden_columns'))


def validate(staging, paths, views, settings):
    """Return no value or raise one grouped, value-free preflight error."""
    ignored = set(_unique(settings.get('ignored_file_columns', []), 'ignored_file_columns'))
    dictionary_views = {normalize(view): value for view, value in settings['dictionary_views'].items()}
    if len(dictionary_views) != len(settings['dictionary_views']):
        raise ValueError('Dictionary contains duplicate normalized view names')
    required_all = set(_unique((settings.get('required_columns') or {}).get('all_files', []),
                               'required all_files columns'))
    required_by_view = {
        normalize(view): set(_unique(columns, f'required columns for {view!r}'))
        for view, columns in ((settings.get('required_columns') or {}).get('by_view') or {}).items()
    }
    if not set(required_by_view) <= set(dictionary_views):
        unknown = sorted(set(required_by_view) - set(dictionary_views))
        raise ValueError(f'Required-column configuration names unknown dictionary views: {unknown}')
    forbidden = load_policy(settings['phi_pii_columns_file'])
    conflicts = (ignored | required_all | set().union(*required_by_view.values(), set())) & forbidden
    if conflicts:
        raise ValueError(f'ALL ALS preflight configuration conflicts with PHI/PII policy: {sorted(conflicts)}')

    failures = []
    for source_id, file_data in staging.items():
        filename = next(iter(file_data))
        view = normalize(views.get(source_id, ''))
        issues = []
        if view not in dictionary_views:
            issues.append(f'no data-dictionary view named {view!r}')
            failures.append((filename, view or '<empty>', issues))
            continue
        path = paths.get(source_id)
        if not path:
            issues.append('file could not be downloaded for header validation')
            failures.append((filename, view, issues))
            continue
        headers = set(read_headers(path))
        forbidden_headers = sorted(headers & forbidden)
        if forbidden_headers:
            issues.append('forbidden PHI/PII headers: ' + ', '.join(forbidden_headers))
        missing_required = sorted((required_all | required_by_view.get(view, set())) - headers)
        if missing_required:
            issues.append('required headers missing: ' + ', '.join(missing_required))
        substantive = headers - ignored
        expected = dictionary_views[view]['fields']
        extras, missing = sorted(substantive - expected), sorted(expected - substantive)
        if extras:
            issues.append('file headers absent from dictionary: ' + ', '.join(extras))
        if missing:
            issues.append('dictionary headers absent from file: ' + ', '.join(missing))
        if issues:
            failures.append((filename, dictionary_views[view]['display_name'], issues))
    if failures:
        lines = ['ALL ALS preflight failed']
        for filename, view, issues in failures:
            lines.append(f'\n{filename} → {view}')
            lines.extend(f'  - {issue}' for issue in issues)
        raise ValueError('\n'.join(lines))
